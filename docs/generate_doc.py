from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2.5)

# Styles
styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)

def add_heading(text, level=1, color=(31, 111, 139)):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        run.font.color.rgb = RGBColor(*color)
        run.font.name = 'Calibri'
    return p

def add_paragraph(text='', bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = 'Calibri'
    return p

def add_bullet(text, bold_part=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_part and text.startswith(bold_part):
        run1 = p.add_run(bold_part)
        run1.bold = True
        run1.font.name = 'Calibri'
        run1.font.size = Pt(11)
        rest = text[len(bold_part):]
        run2 = p.add_run(rest)
        run2.font.name = 'Calibri'
        run2.font.size = Pt(11)
    else:
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)

def add_table(headers, rows, col_widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        hdr[i].paragraphs[0].runs[0].bold = True
        hdr[i].paragraphs[0].runs[0].font.name = 'Calibri'
        hdr[i].paragraphs[0].runs[0].font.size = Pt(10)
        tc = hdr[i]._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), '1F6F8B')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:val'), 'clear')
        tcPr.append(shd)
        for run in hdr[i].paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row in rows:
        r = table.add_row().cells
        for i, val in enumerate(row):
            r[i].text = val
            r[i].paragraphs[0].runs[0].font.name = 'Calibri'
            r[i].paragraphs[0].runs[0].font.size = Pt(10)
    return table

def add_code(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Courier New'
    run.font.size = Pt(9)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), 'F2F2F2')
    shd.set(qn('w:val'), 'clear')
    pPr.append(shd)
    p.paragraph_format.left_indent = Inches(0.3)

# ─── HEADER ───
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r1 = p.add_run('Ddunkin1')
r1.bold = True
r1.font.size = Pt(11)
r1.font.name = 'Calibri'
r2 = p.add_run('                                                                                    IT33')
r2.font.size = Pt(11)
r2.font.name = 'Calibri'

doc.add_paragraph()

# Title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('ShopEase — System Documentation')
run.bold = True
run.font.size = Pt(22)
run.font.name = 'Calibri'
run.font.color.rgb = RGBColor(31, 111, 139)

doc.add_paragraph()

# ─── 1. INTRODUCTION ───
add_heading('1. Introduction')
p = doc.add_paragraph()
r = p.add_run('ShopEase')
r.bold = True
r.font.name = 'Calibri'
r.font.size = Pt(11)
p.add_run(' is a full-featured E-Commerce Web Application built for the course ').font.name = 'Calibri'
r2 = p.add_run('Fundamentals of Database Systems')
r2.italic = True
r2.font.name = 'Calibri'
p.add_run('. It provides a complete online shopping experience — from product browsing and cart management to order placement, delivery tracking, and admin oversight. The system supports multiple user roles including customers, administrators, and delivery riders.').font.name = 'Calibri'

# ─── 2. TECH STACK ───
add_heading('2. Tech Stack')

add_heading('Frontend', level=2)
add_bullet('Framework: React 18 (with Vite)', 'Framework:')
add_bullet('Routing: React Router v6', 'Routing:')
add_bullet('Styling: Tailwind CSS', 'Styling:')
add_bullet('Icons: Lucide React', 'Icons:')
add_bullet('HTTP Client: Axios', 'HTTP Client:')
add_bullet('Notifications: React Hot Toast', 'Notifications:')

add_heading('Backend', level=2)
add_bullet('Framework: Laravel 11 (API-only)', 'Framework:')
add_bullet('Auth: Laravel Sanctum (token-based SPA auth)', 'Auth:')
add_bullet('ORM: Eloquent', 'ORM:')
add_bullet('Language: PHP 8.4', 'Language:')

add_heading('Database', level=2)
add_bullet('System: MySQL', 'System:')
add_bullet('Tool: Navicat / Railway MySQL', 'Tool:')

add_heading('Deployment', level=2)
add_bullet('Platform: Railway', 'Platform:')
add_bullet('Server: FrankenPHP (via Railpack)', 'Server:')
add_bullet('Repository: GitHub', 'Repository:')

# ─── 3. USER ROLES ───
add_heading('3. User Roles & Permissions')
add_table(
    ['Role', 'Description', 'Key Permissions'],
    [
        ['Admin', 'System Overseer', 'Manage products, categories, orders, users, assign riders, view dashboard stats'],
        ['Rider', 'Delivery Personnel', 'View assigned deliveries, update delivery status'],
        ['Customer', 'Shopper', 'Browse products, manage cart, place orders, write reviews, track orders'],
    ]
)

# ─── 4. WORKFLOWS ───
add_heading('4. System Workflows')
add_heading('Shopping Flow', level=2)
for i, step in enumerate([
    'Customer browses products (filterable by category, sortable by price/name).',
    'Customer adds items to cart or uses Buy Now.',
    'Customer proceeds to checkout — enters shipping address and payment method.',
    'Order is placed (Status: pending), payment record created.',
    'Admin processes the order (Status: processing).',
    'Admin assigns a rider for delivery.',
    'Rider updates delivery status: assigned → picked_up → delivered.',
    'Order marked completed upon delivery.',
], 1):
    p = doc.add_paragraph(style='List Number')
    p.add_run(step).font.name = 'Calibri'

add_heading('Review Flow', level=2)
for step in [
    'Customer must have a completed order containing the product.',
    'Customer submits a star rating (1–5) and optional comment.',
    'Review is displayed on the product detail page with average rating.',
]:
    p = doc.add_paragraph(style='List Number')
    p.add_run(step).font.name = 'Calibri'

add_heading('Admin Flow', level=2)
for step in [
    'Admin logs in via /admin panel (dark-themed, separate from customer UI).',
    'Admin manages: Products (CRUD), Categories (inline edit), Orders (status + rider assignment), Users (role management).',
    'Dashboard shows: Total revenue, orders, products, users, monthly revenue chart, low stock alerts, recent orders.',
]:
    p = doc.add_paragraph(style='List Number')
    p.add_run(step).font.name = 'Calibri'

# ─── 5. DATABASE SCHEMA ───
add_heading('5. Database Schema')
add_heading('Tables Overview', level=2)
add_table(
    ['Table', 'Description'],
    [
        ['users', 'All users (customers, admins, riders)'],
        ['categories', 'Product categories'],
        ['products', 'Product listings with stock and pricing'],
        ['carts', 'One cart per user'],
        ['cart_items', 'Items inside a cart'],
        ['orders', 'Placed orders with status and rider info'],
        ['order_items', 'Line items for each order'],
        ['payments', 'Payment record per order'],
        ['reviews', 'Product reviews (one per user per product)'],
        ['personal_access_tokens', 'Sanctum auth tokens'],
    ]
)

add_heading('Key Models', level=2)
add_code("""model User {
  id        Int     @id @autoincrement
  name      String
  email     String  @unique
  password  String
  is_admin  Boolean @default(false)
  is_rider  Boolean @default(false)
  timestamps
}

model Product {
  id          Int      @id @autoincrement
  category_id Int      FK → categories
  name        String
  price       Decimal
  stock       Int
  image       String?
  timestamps
}

model Order {
  id               Int     @id @autoincrement
  user_id          Int     FK → users
  rider_id         Int?    FK → users
  status           Enum    [pending, processing, completed, cancelled]
  delivery_status  Enum    [unassigned, assigned, picked_up, delivered]
  total_amount     Decimal
  shipping_address Text
  timestamps
}

model Review {
  id         Int  @id @autoincrement
  user_id    Int  FK → users
  product_id Int  FK → products
  rating     Int  (1–5)
  comment    Text?
  unique [user_id, product_id]
}""")

# ─── 6. CORE FEATURES ───
add_heading('6. Core Features')
features = [
    ('Product Catalog', '— Category filtering, price/name sorting, search, stock badges, sold count'),
    ('Shopping Cart', '— Add/remove items, quantity stepper, select items for checkout, subtotal'),
    ('Checkout', '— Shipping address, payment method (Cash / GCash / Card)'),
    ('Order Tracking', '— 5-step delivery progress tracker per order'),
    ('Order Cancellation', '— Cancel pending orders with automatic stock restoration'),
    ('Product Reviews', '— Star ratings, eligibility check (must have purchased), average rating on cards'),
    ('Rider System', '— Admin assigns riders; riders update delivery status via their dashboard'),
    ('Admin Dashboard', '— Revenue chart, stat cards, low stock alerts, recent orders'),
    ('Admin Products CRUD', '— Create, edit, delete products with image URL support'),
    ('Admin Categories', '— Inline edit with slug auto-generation'),
    ('Admin User Management', '— View all users, grant/revoke admin access'),
    ('Stock Management', '— Validated before order, decremented on purchase, restored on cancellation'),
    ('Profile Page', '— Edit name/email, change password'),
]
for label, desc in features:
    p = doc.add_paragraph(style='List Bullet')
    r1 = p.add_run(label)
    r1.bold = True
    r1.font.name = 'Calibri'
    r1.font.size = Pt(11)
    r2 = p.add_run(' ' + desc)
    r2.font.name = 'Calibri'
    r2.font.size = Pt(11)

# ─── 7. API ROUTES ───
add_heading('7. API Routes Summary')

add_heading('Public', level=2)
add_table(
    ['Method', 'Endpoint', 'Description'],
    [
        ['POST', '/api/register', 'Register new user'],
        ['POST', '/api/login', 'Login'],
        ['GET', '/api/products', 'List all products'],
        ['GET', '/api/products/{id}', 'Product detail'],
        ['GET', '/api/categories', 'List categories'],
        ['GET', '/api/products/{id}/reviews', 'Product reviews'],
    ]
)
doc.add_paragraph()

add_heading('Authenticated (Customer)', level=2)
add_table(
    ['Method', 'Endpoint', 'Description'],
    [
        ['GET/POST', '/api/cart', 'View / add to cart'],
        ['PATCH', '/api/cart/{item}', 'Update quantity'],
        ['DELETE', '/api/cart/{item}', 'Remove item'],
        ['POST', '/api/orders', 'Place order'],
        ['GET', '/api/orders', 'My orders'],
        ['GET', '/api/orders/{id}', 'Order detail'],
        ['PATCH', '/api/orders/{id}/cancel', 'Cancel order'],
        ['POST', '/api/products/{id}/reviews', 'Submit review'],
        ['PATCH', '/api/profile', 'Update profile'],
    ]
)
doc.add_paragraph()

add_heading('Admin', level=2)
add_table(
    ['Method', 'Endpoint', 'Description'],
    [
        ['GET', '/api/admin/stats', 'Dashboard statistics'],
        ['GET/POST', '/api/admin/products', 'List / create products'],
        ['PUT/DELETE', '/api/admin/products/{id}', 'Update / delete product'],
        ['GET', '/api/admin/orders', 'All orders (filterable by status)'],
        ['PATCH', '/api/admin/orders/{id}/status', 'Update order status'],
        ['PATCH', '/api/admin/orders/{id}/assign-rider', 'Assign rider to order'],
        ['GET', '/api/admin/users', 'All users'],
        ['PATCH', '/api/admin/users/{id}/toggle-admin', 'Grant/revoke admin'],
        ['GET/POST', '/api/admin/riders', 'List / create riders'],
    ]
)
doc.add_paragraph()

add_heading('Rider', level=2)
add_table(
    ['Method', 'Endpoint', 'Description'],
    [
        ['GET', '/api/rider/orders', 'Active deliveries'],
        ['PATCH', '/api/rider/orders/{id}/delivery-status', 'Update delivery status'],
    ]
)

# ─── 8. PAGES ───
add_heading('8. Pages')
add_table(
    ['Page', 'Route', 'Access'],
    [
        ['Home', '/', 'Public'],
        ['Products', '/products', 'Public'],
        ['Product Detail', '/products/:id', 'Public'],
        ['Cart', '/cart', 'Customer'],
        ['Checkout', '/checkout', 'Customer'],
        ['Order Success', '/order-success/:id', 'Customer'],
        ['My Orders', '/orders', 'Customer'],
        ['Order Detail', '/orders/:id', 'Customer'],
        ['Profile', '/profile', 'Customer'],
        ['Login', '/login', 'Public'],
        ['Register', '/register', 'Public'],
        ['Admin Dashboard', '/admin', 'Admin'],
        ['Admin Products', '/admin/products', 'Admin'],
        ['Admin Categories', '/admin/categories', 'Admin'],
        ['Admin Orders', '/admin/orders', 'Admin'],
        ['Admin Users', '/admin/users', 'Admin'],
        ['Admin Riders', '/admin/riders', 'Admin'],
        ['Rider Dashboard', '/rider', 'Rider'],
    ]
)

# ─── 9. DEPLOYMENT ───
add_heading('9. Deployment')
add_bullet('Platform: Railway', 'Platform:')
add_bullet('Runtime: PHP 8.4 + FrankenPHP + MySQL', 'Runtime:')
add_bullet('GitHub: https://github.com/Ddunkin1/e-commerce', 'GitHub:')
add_bullet('Live URL: https://e-commerce-dbms.up.railway.app', 'Live URL:')

doc.save('/home/stubbyjames/Documents/enterprise-dbms/enterprise-dbms/docs/SystemDocumentation.docx')
print('Done!')
